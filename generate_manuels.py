"""
Génère les manuels de prise en main (Word) dans docs/ :
  - Manuel_Enseignant_MECA_AUTO.docx
  - Manuel_Magasinier_MECA_AUTO.docx
  - Fiche_Eleve_MECA_AUTO.docx   (1 page, à imprimer)
Usage : python generate_manuels.py
"""
import re
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BLEU = RGBColor(0x1E, 0x3A, 0x5F)
BLEU_CLAIR = RGBColor(0x2E, 0x6D, 0xA4)
GRIS = RGBColor(0x55, 0x55, 0x55)

ENCADRES = {
    'info':      ('E3F2FD', '90CAF9', 'ℹ️ '),
    'astuce':    ('E8F5E9', 'A5D6A7', '💡 '),
    'attention': ('FFF3E0', 'FFB74D', '⚠️ '),
}


# ── Outils de mise en forme ─────────────────────────────────────

def shade(cell, fill, border=None):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    tcPr.append(shd)
    if border:
        borders = OxmlElement('w:tcBorders')
        for side in ('top', 'left', 'bottom', 'right'):
            b = OxmlElement(f'w:{side}')
            b.set(qn('w:val'), 'single')
            b.set(qn('w:sz'), '6')
            b.set(qn('w:color'), border)
            borders.append(b)
        tcPr.append(borders)


def add_runs(p, text, size=None, color=None):
    """Ajoute du texte avec **gras** à un paragraphe."""
    for i, part in enumerate(re.split(r'\*\*(.+?)\*\*', text)):
        if not part:
            continue
        r = p.add_run(part)
        r.bold = i % 2 == 1
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
    return p


def new_doc(marges=2.0):
    doc = Document()
    st = doc.styles['Normal']
    st.font.name = 'Calibri'
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    st.font.size = Pt(11)
    st.paragraph_format.space_after = Pt(4)
    for lvl, size, color in ((1, 17, BLEU), (2, 13, BLEU_CLAIR), (3, 11.5, BLEU)):
        h = doc.styles[f'Heading {lvl}']
        h.font.name = 'Calibri'
        h.element.rPr.rFonts.set(qn('w:ascii'), 'Calibri')
        h.element.rPr.rFonts.set(qn('w:hAnsi'), 'Calibri')
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.paragraph_format.space_before = Pt(14 if lvl == 1 else 10)
        h.paragraph_format.space_after = Pt(4)
        h.paragraph_format.keep_with_next = True
    for s in doc.sections:
        s.top_margin = s.bottom_margin = Cm(marges)
        s.left_margin = s.right_margin = Cm(marges)
    return doc


def pied_de_page(doc, texte):
    p = doc.sections[0].footer.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(texte + ' — page ')
    r.font.size = Pt(8)
    r.font.color.rgb = GRIS
    for kind, txt in (('begin', None), (None, 'PAGE'), ('end', None)):
        run = p.add_run()
        run.font.size = Pt(8)
        run.font.color.rgb = GRIS
        if kind:
            fc = OxmlElement('w:fldChar')
            fc.set(qn('w:fldCharType'), kind)
            run._r.append(fc)
        else:
            it = OxmlElement('w:instrText')
            it.set(qn('xml:space'), 'preserve')
            it.text = txt
            run._r.append(it)


def bandeau(doc, titre, sous_titre):
    t = doc.add_table(rows=1, cols=1)
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    c = t.cell(0, 0)
    shade(c, '1E3A5F')
    p = c.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(10)
    r = p.add_run('🔧 MECA AUTO')
    r.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(0xBB, 0xDE, 0xFB)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(titre)
    r.bold = True
    r.font.size = Pt(22)
    r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p = c.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run(sous_titre)
    r.font.size = Pt(11)
    r.font.color.rgb = RGBColor(0xE3, 0xF2, 0xFD)
    doc.add_paragraph()


def h(doc, texte, lvl=1):
    return doc.add_heading(texte, lvl)


def para(doc, texte, size=None):
    return add_runs(doc.add_paragraph(), texte, size)


def puces(doc, items, style='List Bullet', size=None):
    for it in items:
        p = doc.add_paragraph(style=style)
        p.paragraph_format.space_after = Pt(2)
        add_runs(p, it, size)


def etapes(doc, items, size=None):
    """Liste numérotée indépendante (redémarre à 1 à chaque appel)."""
    for n, it in enumerate(items, 1):
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.8)
        p.paragraph_format.first_line_indent = Cm(-0.6)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f'{n}.  ')
        r.bold = True
        r.font.color.rgb = BLEU_CLAIR
        if size:
            r.font.size = Pt(size)
        add_runs(p, it, size)


def encadre(doc, kind, texte, size=None):
    fill, border, icone = ENCADRES[kind]
    t = doc.add_table(rows=1, cols=1)
    c = t.cell(0, 0)
    shade(c, fill, border)
    lignes = texte if isinstance(texte, list) else [texte]
    p = c.paragraphs[0]
    add_runs(p, icone + lignes[0], size)
    for l in lignes[1:]:
        add_runs(c.add_paragraph(), l, size)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def tableau(doc, entetes, lignes, largeurs=None, size=10):
    t = doc.add_table(rows=1, cols=len(entetes))
    t.style = 'Table Grid'
    for i, e in enumerate(entetes):
        c = t.rows[0].cells[i]
        shade(c, '1E3A5F')
        r = c.paragraphs[0].add_run(e)
        r.bold = True
        r.font.size = Pt(size)
        r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    for n, ligne in enumerate(lignes):
        cells = t.add_row().cells
        for i, val in enumerate(ligne):
            add_runs(cells[i].paragraphs[0], val, size)
            if n % 2:
                shade(cells[i], 'F5F8FC')
    if largeurs:
        for row in t.rows:
            for i, w in enumerate(largeurs):
                row.cells[i].width = Cm(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)
    return t


# ── Contenus communs ───────────────────────────────────────────

def section_connexion(doc, role):
    h(doc, '1. Se connecter')
    etapes(doc, [
        "Ouvrir le navigateur (Firefox, Chrome ou Edge) et saisir **l'adresse communiquée par l'établissement**, "
        "qui commence par **https://**.",
        "Saisir son **identifiant** (ou son email) et son **mot de passe**, puis **Se connecter**.",
        "À la **première connexion**, le logiciel demande de **choisir un nouveau mot de passe** personnel.",
    ])
    encadre(doc, 'info', [
        "Si le navigateur affiche « **Connexion non sécurisée** » ou « **Risque probable de sécurité** » :",
        "cliquer sur **Paramètres avancés** puis **Continuer vers le site** (ou « Accepter le risque »). "
        "La connexion est bien chiffrée : le certificat est simplement délivré par le serveur du lycée.",
    ])
    puces(doc, [
        "**Changer son mot de passe** : menu **Mon compte**.",
        "**Mot de passe oublié** : lien sous le formulaire de connexion, ou demander au **DDFPT**.",
        "Après **plusieurs erreurs** de mot de passe, le compte est **bloqué quelques minutes** (sécurité).",
        "Toujours cliquer sur **Déconnexion** (en haut à droite) en quittant un poste partagé.",
    ])


def section_statuts(doc):
    h(doc, "Le cycle de vie d'un ordre de réparation (OR)")
    tableau(doc, ['Statut', 'Signification', 'Qui le fait ?'], [
        ['**Ouvert**', "OR créé, véhicule accueilli, travaux pas encore commencés", 'Enseignant (à la création)'],
        ['**En cours**', 'Les élèves travaillent sur le véhicule', 'Enseignant'],
        ['**Terminé**', 'Travaux finis, état des lieux de sortie fait, prêt à facturer', 'Enseignant'],
        ['**Clôturé**', "Facture créée : l'OR est **verrouillé**", 'Automatique à la facturation (magasinier)'],
    ], largeurs=[2.8, 8.2, 5.5])


def section_faq_commune():
    return [
        ("« Formulaire expiré, merci de réessayer »",
         "La page était ouverte depuis longtemps ou la session a changé. **Recharger la page (F5)** et recommencer."),
        ("« Merci de vous connecter pour accéder à cette page »",
         "La session a expiré. Se reconnecter."),
        ("Je ne peux plus modifier un OR",
         "Il est **clôturé** (facturé). Seuls le **DDFPT** et le **magasinier** peuvent encore le modifier."),
    ]


def faq(doc, questions):
    for q, r in questions:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.keep_with_next = True
        rr = p.add_run('❓ ' + q)
        rr.bold = True
        rr.font.color.rgb = BLEU
        p = add_runs(doc.add_paragraph(), '→ ' + r)
        p.paragraph_format.left_indent = Cm(0.6)


# ── Manuel enseignant ──────────────────────────────────────────

def manuel_enseignant():
    doc = new_doc()
    pied_de_page(doc, 'MECA AUTO — Manuel enseignant')
    bandeau(doc, 'Manuel enseignant', "Prise en main du logiciel de gestion de l'atelier")

    para(doc, "Ce manuel explique comment **accueillir un véhicule**, **organiser le travail des élèves** "
              "et **suivre un ordre de réparation (OR)** jusqu'à sa remise au magasinier pour facturation.")
    h(doc, 'En résumé : votre rôle', 2)
    puces(doc, [
        "**Créer les OR** et réaliser les **états des lieux** d'entrée et de sortie avec le client.",
        "**Affecter le travail** aux élèves et **déclarer les sessions de travail** en atelier.",
        "**Valider les incidents** déclarés par les élèves.",
        "Passer l'OR en **Terminé** : le magasinier prend le relais pour la facture.",
        "Gérer **vos classes et vos élèves** (comptes, import Pronote).",
    ])

    section_connexion(doc, 'enseignant')

    h(doc, '2. Le menu')
    tableau(doc, ['Menu', 'À quoi il sert'], [
        ['**Tableau de bord**', "Rendez-vous des 7 prochains jours, alertes, véhicules présents à l'atelier, derniers OR"],
        ['**Ordres de réparation**', 'Liste, recherche et création des OR'],
        ['**Clients** / **Véhicules**', 'Fiches clients et véhicules, historique'],
        ['**📅 Agenda**', 'Rendez-vous clients'],
        ['**👨‍🎓 Élèves**', 'Vos classes, comptes élèves, import CSV'],
        ['**✅ Checklist**', 'Points du contrôle qualité proposés lors des états des lieux'],
        ['**Mon compte**', 'Changer son mot de passe'],
    ], largeurs=[4.5, 12])
    encadre(doc, 'astuce', "Sur **tablette ou smartphone**, le menu est replié derrière le bouton **☰ Menu**. "
                           "Un bouton **🌙 Mode sombre** est disponible en haut de page.")

    h(doc, '3. Accueillir un véhicule : créer un OR')
    para(doc, "Menu **Ordres de réparation** → **Nouvel OR**.")
    etapes(doc, [
        "**Immatriculation** : si le véhicule est déjà connu, ses informations et celles du client se remplissent seules. "
        "Sinon, compléter **marque, modèle, année**, et si possible **VIN** et **énergie**.",
        "Cocher **CT valide** et **Assurance valide** après vérification des documents.",
        "**Client** : nom, prénom et **téléphone à 10 chiffres (obligatoire)**. L'**email** est facultatif mais "
        "permet d'envoyer au client les avis de suivi et la facture.",
        "**Travaux demandés** : décrire la demande du client.",
        "Choisir la **classe** et l'**élève responsable**.",
        "Facultatif : cocher **📅 Prise de rendez-vous** pour planifier l'intervention dans l'agenda.",
        "Cocher **🚗 Faire l'état des lieux d'entrée après création** (recommandé), puis **enregistrer**.",
    ])
    encadre(doc, 'attention', "L'OR reçoit un **numéro unique**. Il est au statut **Ouvert**.")
    section_statuts(doc)

    h(doc, "4. Les états des lieux (entrée et sortie)")
    para(doc, "Ils se font **avec le client**, depuis la fiche de l'OR, boutons **➕ État des lieux ENTRÉE** "
              "et **➕ État des lieux SORTIE**. Les élèves peuvent aussi les remplir.")
    puces(doc, [
        "**Kilométrage** (obligatoire) et **niveau de carburant**.",
        "**Dommages existants** (carrosserie, vitres, intérieur…) et **observations**.",
        "**📦 Inventaire des objets présents** dans le véhicule, en texte libre, puis case **signée** par le client. "
        "Cette étape protège l'atelier et les élèves en cas de litige (objet disparu…).",
        "**Checklist contrôle qualité** : facultative.",
    ])
    puces(doc, [
        "À la **sortie**, le kilométrage est **pré-rempli** avec celui de l'entrée : il suffit de l'ajuster.",
        "Boutons **🖨️ Imprimer** pour faire **signer le document papier** au client.",
    ])
    encadre(doc, 'attention', "L'état des lieux de **SORTIE** est **obligatoire** avant de clôturer l'OR.")

    h(doc, '5. Pendant les travaux')
    h(doc, "Faire avancer le statut", 2)
    para(doc, "Sur la fiche de l'OR, encadré **Modifier le statut** : passer l'OR en **En cours** "
              "dès que les élèves commencent.")

    h(doc, 'Les interventions des élèves', 2)
    para(doc, "Tableau **Interventions élèves** → **Ajouter une intervention** :")
    puces(doc, [
        "choisir l'**élève**, le nombre d'**heures** (par demi-heure) ;",
        "une **fourniture** utilisée si besoin (elle sera **facturée**) ;",
        "une **description** du travail réalisé.",
    ])
    para(doc, "Les boutons ✏️ et 🗑️ permettent de **corriger** ou **supprimer** une ligne tant que l'OR n'est pas clôturé. "
              "Les élèves peuvent aussi saisir eux-mêmes leurs interventions.")
    encadre(doc, 'astuce', "Un élève voit dans sa liste les OR **sur lesquels il a une intervention**. "
                           "Pour qu'un élève retrouve un OR, ajoutez-lui une première intervention.")

    h(doc, 'Les sessions de travail (traçabilité)', 2)
    para(doc, "Bouton **+ Déclarer une session** : à faire à **chaque séance** en atelier sur le véhicule.")
    puces(doc, [
        "**Date et heure**, **classe**, **élèves présents** ;",
        "**zone du véhicule** concernée (moteur, habitacle, train roulant, carrosserie, électricité, autre) ;",
        "**observations** libres. L'enseignant déclarant est enregistré automatiquement.",
    ])
    encadre(doc, 'info', [
        "**✅ Certifier et enregistrer** fige la session : elle ne pourra **plus être modifiée**.",
        "**💾 Enregistrer (sans certifier)** permet de la compléter puis de la certifier plus tard.",
        "Ces sessions indiquent **qui était sur le véhicule, et quand**. C'est essentiel en cas d'incident.",
    ])

    h(doc, 'Véhicule en attente de pièces', 2)
    para(doc, "Bouton **Modifier** de l'OR → cocher **🚧 En attente de pièces détachées** et ajouter une **remarque**. "
              "Si l'OR est **En cours**, il apparaît dans les **alertes du tableau de bord** avec le nombre de jours d'attente.")

    h(doc, 'Tarif', 2)
    para(doc, "Dans **Modifier** : choisir le **mode tarif**, soit **Forfait** (liste des forfaits), soit **Taux horaire** avec un montant. "
              "Cocher **Pas de facturation** pour des **travaux internes** au lycée.")

    h(doc, '6. Incidents : vol, dégradation, objet manquant')
    para(doc, "Encadré **⚠️ Incidents déclarés** → **Déclarer un incident** : type, date et heure du constat, "
              "**description précise**, objet(s) concerné(s).")
    tableau(doc, ['Déclaré par', 'Ce qui se passe'], [
        ['**Enseignant**', 'Incident **enregistré et validé** immédiatement'],
        ['**Élève**', "Incident **en attente** : un enseignant doit cliquer sur **✅ Valider** sur la fiche de l'OR"],
    ], largeurs=[4, 12.5])
    para(doc, "Une fois validé, seul le **DDFPT** peut modifier un incident, et chaque modification est tracée. "
              "Le DDFPT peut recevoir une **alerte par email** si cette option est activée.")

    h(doc, '7. Dépollution')
    para(doc, "Encadré **♻️ Dépollution** : cocher si le client **récupère ses pièces usagées** et/ou **ses fluides** "
              "(huile, liquide de frein…). Si ce n'est **pas** le cas, des **frais de dépollution** sont ajoutés "
              "automatiquement à la facture. Cliquer sur **💾 Enregistrer**.")

    h(doc, '8. Autres outils de la fiche OR')
    puces(doc, [
        "**Contrôle visuel** : grille de contrôle du véhicule (bon / mauvais), imprimable.",
        "**Imprimer OR** : l'ordre de réparation au format papier.",
        "**🚗 Document véhicule** : fiche à poser derrière le pare-brise.",
        "**Rendez-vous liés** : les RDV du client, avec **+ Nouveau RDV**.",
    ])

    h(doc, "9. Terminer l'OR")
    etapes(doc, [
        "Vérifier que les **interventions** et les **sessions** sont complètes.",
        "Réaliser l'**état des lieux de SORTIE** avec le client.",
        "Passer le statut en **Terminé**.",
        "Le **magasinier** crée la facture : l'OR passe automatiquement en **Clôturé**.",
    ])
    encadre(doc, 'attention', "Un OR **clôturé** est en **lecture seule** pour les enseignants, "
                              "y compris les sessions de travail et les incidents.")

    h(doc, '10. Gérer ses classes et ses élèves')
    para(doc, "Menu **👨‍🎓 Élèves**.")
    puces(doc, [
        "**Classes** : créer une classe (ex. 1MAVA), la renommer, l'activer ou la désactiver.",
        "**Ajouter un élève** : nom, prénom, login, mot de passe, classe, email.",
        "**Modifier un élève** : bouton ✏️. Pour **réinitialiser son mot de passe**, saisir le nouveau "
        "(laisser vide = inchangé).",
    ])
    h(doc, 'Importer une classe depuis Pronote (CSV)', 2)
    etapes(doc, [
        "Télécharger le **modèle CSV** proposé sur la page.",
        "Colonnes : **nom ; prénom ; classe ; date de naissance ; adresse e-mail** "
        "(séparateur « ; » ou « , », détecté automatiquement).",
        "Cliquer sur **📥 Importer CSV**. Les classes inexistantes sont créées automatiquement.",
    ])
    tableau(doc, ['Élément', 'Règle'], [
        ['**Identifiant** généré', '3 premières lettres du prénom + 3 premières lettres du nom, en minuscules (ex. Martin Dupont → **mardup**)'],
        ['**Mot de passe** provisoire', 'Date de naissance **JJMMAAAA** (ex. 15032008), ou l\'identifiant si pas de date'],
    ], largeurs=[4.5, 12])
    encadre(doc, 'attention', "Communiquez les identifiants **individuellement** et demandez aux élèves de changer "
                              "leur mot de passe dès la première connexion.")

    h(doc, '11. Checklist contrôle qualité')
    para(doc, "Menu **✅ Checklist** : ajouter, modifier, activer ou désactiver les points proposés lors des états des lieux.")

    h(doc, '12. Agenda')
    para(doc, "Menu **📅 Agenda** → **Nouveau RDV** : client, véhicule, titre, date, heure et durée. "
              "Les RDV peuvent être modifiés ou supprimés (une confirmation est demandée). "
              "Le tableau de bord affiche les RDV **d'aujourd'hui et des 7 jours suivants**.")

    h(doc, 'Questions fréquentes')
    faq(doc, section_faq_commune() + [
        ("Un élève ne voit pas l'OR sur lequel il travaille",
         "Lui ajouter une **intervention** sur cet OR (voir partie 5)."),
        ("Un élève a oublié son mot de passe",
         "Menu **Élèves** → ✏️ sur l'élève → saisir un **nouveau mot de passe**."),
        ("Je ne peux pas clôturer l'OR",
         "Il manque l'**état des lieux de sortie**. La clôture se fait normalement par la **facturation**."),
    ])
    return doc


# ── Manuel magasinier ──────────────────────────────────────────

def manuel_magasinier():
    doc = new_doc()
    pied_de_page(doc, 'MECA AUTO — Manuel magasinier')
    bandeau(doc, 'Manuel magasinier', 'Facturation, encaissement, pièces et fournitures')

    para(doc, "Le magasinier **établit les factures**, **suit les paiements**, gère les **petites fournitures** "
              "de l'atelier (scotch, WD-40…) et les **tarifs** (taux horaire, forfaits, frais de dépollution).")
    h(doc, 'En résumé : votre rôle', 2)
    puces(doc, [
        "**Facturer** les OR passés en **Terminé** par les enseignants.",
        "Télécharger ou **envoyer la facture** au client, puis la marquer **Soldée** à l'encaissement.",
        "Décider si les **frais de dépollution** sont **pris en charge par le lycée**.",
        "Tenir à jour le **catalogue des fournitures** et les **tarifs**.",
    ])

    section_connexion(doc, 'magasinier')

    h(doc, '2. Le menu')
    tableau(doc, ['Menu', 'À quoi il sert'], [
        ['**Tableau de bord**', "RDV à venir, alertes, véhicules à l'atelier, derniers OR"],
        ['**Ordres de réparation**', 'Consulter les OR, filtrer par statut (ex. **Terminé** = à facturer)'],
        ['**Clients** / **Véhicules**', 'Coordonnées des clients (email pour l\'envoi des factures)'],
        ['**📅 Agenda**', 'Rendez-vous clients'],
        ['**Factures**', "Toutes les factures, filtrables par année, avec l'état du paiement"],
        ['**🧴 Fournitures**', 'Fournitures, taux horaire, forfaits, frais de dépollution'],
        ['**Mon compte**', 'Changer son mot de passe'],
    ], largeurs=[4.5, 12])

    section_statuts(doc)

    h(doc, '3. Créer une facture')
    etapes(doc, [
        "Menu **Ordres de réparation** : repérer les OR au statut **Terminé**.",
        "Ouvrir l'OR et **vérifier** : mode tarif et montant, fournitures saisies dans les interventions, "
        "encadré **♻️ Dépollution**.",
        "Cliquer sur **Créer la Facture**, puis **confirmer**.",
    ])
    para(doc, "Le montant est **calculé automatiquement** :")
    tableau(doc, ['Élément', 'Origine'], [
        ['**Travaux**', "Forfait ou taux horaire choisi sur l'OR"],
        ['**+ Fournitures**', 'Fournitures saisies dans les interventions des élèves (prix × quantité)'],
        ['**+ Dépollution**', 'Si le client ne récupère pas ses pièces et/ou ses fluides (sauf prise en charge)'],
    ], largeurs=[4, 12.5])
    encadre(doc, 'attention', [
        "La création de la facture **clôture l'OR** : il devient **en lecture seule** pour les enseignants.",
        "Vous et le DDFPT pouvez encore le modifier. Seul le **DDFPT** peut **supprimer** un OR facturé.",
    ])
    encadre(doc, 'astuce', "Pour corriger un montant, faites-le **avant** de créer la facture, "
                           "avec le bouton **Modifier** de l'OR (mode tarif, forfait, montant).")

    h(doc, '4. Frais de dépollution')
    para(doc, "Sur la fiche de l'OR, encadré **♻️ Dépollution** :")
    puces(doc, [
        "l'enseignant indique si le client **récupère ses pièces** et/ou **ses fluides** ;",
        "**vous seul** (avec le DDFPT) pouvez cocher **🏫 Frais de dépollution pris en charge par le lycée** : "
        "les frais passent à **0 €** ;",
        "cliquer sur **💾 Enregistrer** avant de créer la facture.",
    ])

    h(doc, '5. Après la facture : envoi et paiement')
    para(doc, "Menu **Factures** → ouvrir la facture :")
    puces(doc, [
        "**⬇️ Télécharger PDF** : pour l'imprimer ou la remettre au client. Le PDF contient les informations "
        "de l'établissement, le logo et les **conditions de service** en bas de page.",
         "**✉ Envoyer par mail** : envoie le PDF au client (si son **email** est renseigné dans sa fiche).",
        "**Statut de paiement** : **⏳ En attente de paiement** → **✅ Soldée** une fois l'argent encaissé. "
        "On peut revenir en arrière en cas d'erreur.",
    ])
    encadre(doc, 'info', "Chaque changement de statut de paiement est **enregistré dans le journal** "
                         "(qui, quand), pour la traçabilité de la caisse.")

    h(doc, '6. Fournitures et tarifs')
    para(doc, "Menu **🧴 Fournitures** :")
    tableau(doc, ['Rubrique', 'Ce que vous pouvez faire'], [
        ['**Fournitures**', "Ajouter (nom + prix unitaire), modifier directement dans le tableau, supprimer. "
                            "Ce sont les articles que les élèves sélectionnent dans leurs interventions."],
        ['**⏱️ Taux horaire**', "Prix de l'heure de main-d'œuvre, utilisé pour les OR « au taux horaire »."],
        ['**💰 Forfaits**', 'Prestations à prix fixe (vidange, freinage…) : ajouter, modifier, supprimer.'],
        ['**♻️ Frais de dépollution**', 'Montants appliqués si pièces et/ou fluides ne sont pas récupérés.'],
    ], largeurs=[4.5, 12])
    encadre(doc, 'attention', "Une modification de tarif s'applique aux **prochaines** factures. "
                              "Les factures déjà créées ne changent pas.")

    h(doc, '7. Autres actions possibles')
    puces(doc, [
        "**Créer et modifier des OR**, des **clients** et des **véhicules**.",
        "**Déclarer un incident** sur un OR.",
        "Consulter l'**agenda** et les **états des lieux** (impression).",
    ])

    h(doc, 'Questions fréquentes')
    faq(doc, section_faq_commune() + [
        ("Le bouton « Créer la Facture » n'apparaît pas",
         "L'OR n'est pas encore **Terminé**, ou une facture **existe déjà** (bouton **Voir la Facture**)."),
        ("Le bouton « Envoyer par mail » n'apparaît pas",
         "Le client n'a pas d'**email** dans sa fiche : le compléter dans **Clients**. Ou la facture a déjà été envoyée."),
        ("Le client ne paie pas tout de suite",
         "Laisser la facture **En attente de paiement**. La liste **Factures** affiche l'état de paiement de chacune."),
    ])
    return doc


# ── Fiche élève (1 page) ───────────────────────────────────────

def fiche_eleve():
    doc = new_doc(marges=1.3)
    st = doc.styles['Normal']
    st.font.size = Pt(10)
    st.paragraph_format.space_after = Pt(1)
    for lvl in (1, 2):
        hs = doc.styles[f'Heading {lvl}']
        hs.paragraph_format.space_before = Pt(6)
        hs.paragraph_format.space_after = Pt(2)
    doc.styles['Heading 1'].font.size = Pt(12.5)
    pied_de_page(doc, 'MECA AUTO — Fiche élève')
    bandeau(doc, 'Fiche élève', "L'essentiel pour bien démarrer")
    doc.paragraphs[-1].paragraph_format.space_after = Pt(0)

    S = 9.5
    h(doc, '🔑 Me connecter')
    etapes(doc, [
        "Adresse donnée par l'enseignant, commence par **https://**. Si le navigateur affiche un avertissement : "
        "**Paramètres avancés** → **Continuer vers le site**.",
        "**Identifiant** = 3 lettres du prénom + 3 lettres du nom (ex. Martin Dupont → **mardup**).",
        "**Mot de passe** de départ = ma date de naissance **JJMMAAAA**. Je le **change** à la première connexion.",
    ], size=S)
    encadre(doc, 'attention', "Mon mot de passe est **personnel** : je ne le donne à personne. "
                              "Je clique sur **Déconnexion** en quittant le poste.", size=S)

    h(doc, "🧭 Ce que je vois")
    tableau(doc, ['Menu', 'Contenu'], [
        ['**Tableau de bord**', 'Mes OR en cours'],
        ['**Ordres de réparation**', 'Les OR sur lesquels je travaille (les coordonnées du client sont masquées)'],
        ['**📋 Mes interventions**', 'Tout mon travail : dates, véhicules, heures, total de mes heures'],
    ], largeurs=[4.8, 13.5], size=S)

    h(doc, '🔧 Ce que je fais sur un OR')
    tableau(doc, ['Action', 'Où ?', 'Comment'], [
        ['**Saisir mon travail**', 'Ajouter une intervention',
         "Heures (par demi-heure), fourniture utilisée, **description claire** de ce que j'ai fait"],
        ['**État des lieux**', '➕ ENTRÉE / ➕ SORTIE',
         'Kilométrage, carburant, dommages, **inventaire des objets** présents dans le véhicule'],
        ['**Contrôle visuel**', 'Bouton Contrôle visuel', 'Cocher chaque point : ✓ Bon ou ✗ Mauvais'],
        ['**Signaler un problème**', '⚠️ Déclarer un incident',
         'Vol, dégradation, objet manquant, anomalie : **je décris précisément** ce que j\'ai constaté'],
    ], largeurs=[3.6, 4.2, 10.5], size=S)
    encadre(doc, 'info', "Mes incidents sont **validés par l'enseignant** avant d'être enregistrés. "
                         "Signaler un problème, c'est se **protéger** : ce n'est jamais une faute.", size=S)

    h(doc, '✅ Les bons réflexes')
    puces(doc, [
        "Je saisis mes interventions **le jour même**, avec des mots précis (pièce, opération, résultat).",
        "Je note **tout objet** présent dans le véhicule à l'état des lieux d'entrée.",
        "Je signale **immédiatement** toute casse ou disparition à mon enseignant **et** dans le logiciel.",
        "Un OR **Clôturé** ne peut plus être modifié : je vérifie ma saisie avant la fin des travaux.",
        "Message « **Formulaire expiré** » : je **recharge la page (F5)** et je recommence.",
        "Mot de passe oublié : je demande à **mon enseignant**.",
    ], size=S)
    return doc


if __name__ == '__main__':
    import os
    os.makedirs('docs', exist_ok=True)
    for nom, fn in (('Manuel_Enseignant_MECA_AUTO.docx', manuel_enseignant),
                    ('Manuel_Magasinier_MECA_AUTO.docx', manuel_magasinier),
                    ('Fiche_Eleve_MECA_AUTO.docx', fiche_eleve)):
        fn().save(os.path.join('docs', nom))
        print('✓', nom)
